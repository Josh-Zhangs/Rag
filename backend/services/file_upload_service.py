import logging
import pdfplumber
import pytesseract
import os
import sys
import io
import re
import threading
import zipfile
import shutil
import subprocess
import xml.etree.ElementTree as ET
from datetime import datetime
import pandas as pd
from typing import List, Dict, Tuple, Optional, Any
from PIL import Image, ImageOps
from concurrent.futures import ThreadPoolExecutor, TimeoutError, as_completed
from unicodedata import normalize

logger = logging.getLogger(__name__)

# =========================
# 1. OCR 配置
# =========================
# 获取 backend/ 根目录，兼容开发模式与 PyInstaller 打包模式
if getattr(sys, 'frozen', False):
    BASE_DIR = sys._MEIPASS
else:
    BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Tesseract OCR 可执行文件路径
TESSERACT_EXE_PATH = os.path.join(BASE_DIR, "runtime", "tesseract", "tesseract.exe")
# Tesseract 语言包目录
TESSDATA_DIR = os.path.join(BASE_DIR, "runtime", "tesseract", "tessdata")

# 设置 Tesseract 命令路径
pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE_PATH
# 设置语言包环境变量
os.environ["TESSDATA_PREFIX"] = TESSDATA_DIR

# 支持的语言：日语+英语
SUPPORTED_LANGUAGES = 'jpn+eng'
# OCR 超时时间（秒）
OCR_TIMEOUT = 8
# 最大图像像素
MAX_IMG_PIXEL = 4000
# 页面处理线程数
PAGE_WORKERS = 8
# OCR 处理线程数
OCR_WORKERS = 12
# OCR 图像最大像素限制
MAX_OCR_IMG_PX = 2000

# 兼容不同版本的 PIL 图像重采样方式
try:
    BILINEAR_FILTER = Image.Resampling.BILINEAR
except AttributeError:
    BILINEAR_FILTER = Image.BILINEAR

# ================================================================
# 2. 共通工具
# ================================================================

# --- 1.1 文档元数据包装 ---
# 生成文档元数据头部，PDF/Image/格式共用
def wrap_markdown(file_name: str, content: str, file_size_mb: float,
                  doc_type: str = "PDF") -> str:
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    return (
        f"# {doc_type}ドキュメント: {file_name}\n\n"
        f"## ドキュメント情報\n"
        f"- **元ファイル**: {file_name}\n"
        f"- **ファイルサイズ**: {file_size_mb:.2f} MB\n"
        f"- **変換日時**: {now}\n\n"
        f"---\n\n"
        f"{content}"
    )

# --- 1.2 OCR 工具 ---
# 将预处理后的图片送入 OCR 引擎提取文字
def _ocr_raw(img, lang, psm):
    config = f'--oem 1 --psm {psm} -c preserve_interword_spaces=1'
    try:
        return pytesseract.image_to_string(img, lang=lang, config=config)
    except Exception:
        return ""

# 识别小图/低对比度图片的文字，提高准确率：
# 先使用 PSM 11（稀疏文本），再使用 PSM 3（自动识别）
def _ocr_fast(img, lang):
    text = _ocr_raw(img, lang, 11)
    if text.strip():
        return text
    return _ocr_raw(img, lang, 3)

# 提高小图/低对比度图片的文字识别准确率。
def _preprocess_fast(pil_img: Image.Image) -> Image.Image:
    w, h = pil_img.size
    max_dim = max(w, h)
    if max_dim < 500:
        scale = min(1.5, MAX_IMG_PIXEL / max_dim)
        new_w, new_h = int(w * scale), int(h * scale)
        pil_img = pil_img.resize((new_w, new_h), BILINEAR_FILTER)
    gray = ImageOps.grayscale(pil_img)
    return ImageOps.autocontrast(gray, cutoff=2)

# 清洗 OCR 结果，合并断行的文本，合并相同行的文本。
def _clean_ocr_text(text: str) -> str:
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    cleaned_lines = []
    current_line = ""
    for line in lines:
        if len(line) < 30 and line.endswith(':'):
            if current_line:
                cleaned_lines.append(current_line)
            cleaned_lines.append(line)
            current_line = ""
        elif (line.endswith('.') or
              (len(line) < 50 and line and line[0].isupper())):
            if current_line:
                cleaned_lines.append(current_line)
            current_line = line
        else:
            if current_line:
                current_line += " " + line
            else:
                current_line = line
    if current_line:
        cleaned_lines.append(current_line)
    return '\n'.join(cleaned_lines)

# 从 PIL Image 提取文本，先使用 _ocr_fast 识别，再使用 _preprocess_fast 回退
def _ocr_image_to_text(pil_img: Image.Image, gray: Image.Image) -> str:
    # 第一阶段 OCR
    text = _ocr_fast(gray, SUPPORTED_LANGUAGES)
    if not text.strip():
        # 增强预处理后重试
        enhanced = _preprocess_fast(pil_img)
        # 增强后的图像同样需要缩放，防止过大图片拖慢 OCR
        w, h = enhanced.size
        max_dim = max(w, h)
        if max_dim > MAX_OCR_IMG_PX:
            scale = MAX_OCR_IMG_PX / max_dim
            enhanced = enhanced.resize(
                (int(w * scale), int(h * scale)), BILINEAR_FILTER
            )
        text = _ocr_fast(enhanced, SUPPORTED_LANGUAGES)
    # 清洗 OCR 结果
    return _clean_ocr_text(text)

# 将 PIL Image 转换为灰度图像，支持所有 PIL 颜色模式（L/RGBA/P/CMYK/RGB）
def _convert_to_gray(pil_img: Image.Image) -> Image.Image:
    if pil_img.mode == 'L':
        return pil_img
    elif pil_img.mode in ('RGBA', 'P'):
        return ImageOps.grayscale(pil_img)
    elif pil_img.mode == 'CMYK':
        return ImageOps.grayscale(pil_img.convert('RGB'))
    else:
        return ImageOps.grayscale(pil_img)

# 将 PIL 图片转为灰度、限制尺寸后送入 OCR 引擎提取文字
def _ocr_image_from_pil(pil_img: Image.Image) -> str:
    # 灰度转换（通过共通 helper，覆盖所有 PIL 模式）
    gray = _convert_to_gray(pil_img)

    # 缩放过大图片，避免 OCR 超时
    w, h = gray.size
    max_dim = max(w, h)
    if max_dim > MAX_OCR_IMG_PX:
        scale = MAX_OCR_IMG_PX / max_dim
        gray = gray.resize((int(w * scale), int(h * scale)), BILINEAR_FILTER)

    return _ocr_image_to_text(pil_img, gray)


# --- 1.3 表格工具 ---
# 判断键值对中的键是否有效
def _is_valid_key(key: str) -> bool:
    # 纯数字不是有效键
    if re.match(r'^\d+$', key):
        return False
    # 长度检查
    if len(key) < 1 or len(key) > 100:
        return False
    # 货币符号开头不是键
    if re.match(r'^[R$€£¥]\s*\d', key):
        return False
    # 百分比不是键
    if re.match(r'^\d+(?:\.\d+)?\s*%$', key):
        return False
    # 包含中日文字符是有效键
    if re.search(r'[\u3040-\u30FF\u4E00-\u9FFF]', key):
        return True
    # 首字母大写或日文假名开头是有效键
    if key and (key[0].isupper() or re.match(r'^[\u3040-\u30FF]', key)):
        return True
    return True

# 判断键值对中的值是否有效
def _is_valid_value(value: str) -> bool:
    # 空值无效
    if not value.strip():
        return False
    # 纯特殊字符无效
    if re.match(r'^[^\w]+$', value):
        return False
    return True

# 将 DataFrame 转换为 Markdown 管道表格
def _md_table(df) -> str:
    lines = []
    # 生成表头行
    lines.append("| " + " | ".join(str(c) for c in df.columns) + " |")
    # 生成分隔行
    lines.append("|" + "|".join(["---"] * len(df.columns)) + "|")
    # 生成数据行
    for _, row in df.iterrows():
        lines.append("| " + " | ".join(str(c) for c in row) + " |")
    return "\n".join(lines)

# 从表格中提取键值对
def _extract_kv_from_table(df) -> Dict[str, str]:
    kv_pairs = {}
    # 两列表格：第一列作为键，第二列作为值
    if len(df.columns) == 2:
        for _, row in df.iterrows():
            key = str(row.iloc[0]).strip()
            value = str(row.iloc[1]).strip()
            if (key and value and
                    _is_valid_key(key) and
                    _is_valid_value(value)):
                kv_pairs[key] = value
    # 多列表格：提取财务相关数据
    elif len(df.columns) > 2:
        for _, row in df.iterrows():
            if len(row) >= 2:
                key = str(row.iloc[0]).strip()
                for ci in range(1, min(3, len(row))):
                    value = str(row.iloc[ci]).strip()
                    # 判断是否为财务数值（金额、百分比、倍数等）
                    is_financial = any(
                        re.match(p, value, re.IGNORECASE) for p in [
                            r'^[R$€£¥]\s*[\d,]+(?:\.\d+)?\s*(?:million|m|billion|b|trillion|t)?$',
                            r'^\d+(?:\.\d+)?\s*%$',
                            r'^[\d,]+(?:\.\d+)?$',
                            r'^\d+(?:\.\d+)?\s*x$',
                            r'.*(?:million|billion|trillion|m|b|t)$',
                        ]
                    )
                    if is_financial and _is_valid_key(key):
                        kv_pairs[f"{key} (Col {ci+1})"] = value
    return kv_pairs

# 将二维表格数据转换为 Markdown（共用：PDF/DOCX/XLSX）
def _convert_table_to_md(table_data: List[List], table_idx: int) -> str:
    # 清理表格数据：去除空行和空单元格
    cleaned_data = []
    for row in table_data:
        if row:
            cleaned_row = [
                str(cell).strip() if cell is not None else ""
                for cell in row
            ]
            if any(cell for cell in cleaned_row):
                cleaned_data.append(cleaned_row)
    if not cleaned_data:
        return ""

    # 转换为 DataFrame 便于处理
    df = pd.DataFrame(cleaned_data)
    # 删除全空行
    df = df.dropna(how='all').reset_index(drop=True)
    # 删除全空列
    df = df.loc[:, df.any()].reset_index(drop=True)
    if df.empty:
        return ""

    # 添加表格标题
    header = f"#### Table {table_idx + 1}\n\n"
    body = _md_table(df)

    # 从表格中提取键值对
    kv = _extract_kv_from_table(df)
    if kv:
        kv_lines = ["\n**Key-Value Pairs in Table**:\n"]
        for k, v in kv.items():
            kv_lines.append(f"- **{k}**: {v}\n")
        body += "".join(kv_lines)

    return header + body


# --- 1.4 LibreOffice 格式转换工具 ---
# 获取 soffice.exe 的完整路径
def _get_soffice_path() -> str:
    import platform
    
    # 方法1：检查 PATH（shutil 已在模块顶部导入）
    soffice = shutil.which('soffice')
    if soffice:
        return soffice
    
    # 方法2：检查 Windows 默认安装路径
    if platform.system() == 'Windows':
        default_paths = [
            r'C:\Program Files\LibreOffice\program\soffice.exe',
            r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
        ]
        for path in default_paths:
            if os.path.exists(path):
                return path
    
    return 'soffice'  # fallback

# 检测 LibreOffice 是否已安装并可调用
def is_libreoffice_available() -> bool:
    return _get_soffice_path() != 'soffice' or os.path.exists(
        r'C:\Program Files\LibreOffice\program\soffice.exe'
    ) or os.path.exists(
        r'C:\Program Files (x86)\LibreOffice\program\soffice.exe'
    )

# 共通 LibreOffice 格式转换（.doc→.docx / .xls→.xlsx）
def _convert_with_libreoffice(input_path: str, output_dir: str,
                              target_format: str) -> str:
    soffice_path = _get_soffice_path()

    # 调用 LibreOffice 无头模式进行转换
    result = subprocess.run(
        [
            soffice_path, "--headless", "--convert-to", target_format,
            "--outdir", output_dir, input_path
        ],
        capture_output=True,
        text=True,
        timeout=60  # 60 秒超时
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice 変換に失敗しました: {result.stderr}")

    # 构造输出文件名：原文件名（不含扩展名）+ 目标扩展名
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    converted_path = os.path.join(output_dir, f"{base_name}.{target_format}")
    if not os.path.exists(converted_path):
        raise FileNotFoundError(f"変換後のファイルが見つかりません: {converted_path}")
    return converted_path

# 转换 .doc 为 .docx
def convert_doc_to_docx(doc_path: str, output_dir: str) -> str:
    return _convert_with_libreoffice(doc_path, output_dir, "docx")

# 转换 .xls 为 .xlsx
def convert_xls_to_xlsx(xls_path: str, output_dir: str) -> str:
    return _convert_with_libreoffice(xls_path, output_dir, "xlsx")


# --- 1.5 日文文本后处理 ---
def normalize_jp_text(raw_text: str) -> str:
    # 统一日文全半角、符号、日期、空格
    # Unicode 规范化
    text = normalize("NFKC", raw_text)
    # 全角空格转半角
    text = text.replace("\u3000", " ")
    # 全角括号转半角
    text = text.replace("（", "(").replace("）", ")")
    # 方括号转换
    text = text.replace("【", "[").replace("】", "]")
    # 规范化令和日期格式
    text = re.sub(r"令和\s*([0-9０-９]+)\s*年",
                  lambda m: f"令和{m.group(1)}年", text)
    # 合并多余空行
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def clean_markdown(md: str) -> str:
    # 删除 PDF 自带页码、分割线、纯数字碎片，去除无效噪声
    # 先规范化日文文本
    md = normalize_jp_text(md)
    # 删除 Page 页码
    md = re.sub(r"^\s*Page\s*\d+\s*$", "", md, flags=re.MULTILINE)
    # 删除 1/100 格式页码
    md = re.sub(r"^\s*\d+\s*/\s*\d+\s*$", "", md, flags=re.MULTILINE)
    # 删除分隔线
    md = re.sub(r"^[-=]{3,}$", "", md, flags=re.MULTILINE)
    # 逐行过滤
    raw_lines = md.splitlines()
    valid_lines = []
    for ln in raw_lines:
        s = ln.strip()
        if not s:
            valid_lines.append("")
            continue
        # 过滤纯数字和日期碎片
        if re.fullmatch(r"[0-9０-９\-\/\s]+", s):
            continue
        # 过滤短数字行
        if len(s) <= 2 and re.search(r"\d", s):
            continue
        valid_lines.append(s)
    # 合并结果并清理多余空行
    final = "\n".join(valid_lines)
    final = re.sub(r"\n{3,}", "\n\n", final).strip()
    return final


# ================================================================
# 3. PDF 解析
# ================================================================

class PDFToMarkdownKVConverter:

    def __init__(self, use_ocr: bool = True):
        # 是否启用 OCR 识别图像中的文字
        self.use_ocr = use_ocr
        # 创建 OCR 处理线程池
        self._ocr_executor = ThreadPoolExecutor(max_workers=OCR_WORKERS)
        # PDF 对象访问锁，防止多线程竞争
        self._pdf_lock = threading.Lock()

    # 处理单页 PDF 内容，包括文本、表格、图像等
    def _process_single_page(self, page, page_idx: int) -> str:
        # 初始化页面 Markdown 内容，添加页面标题
        page_md = f"## Page {page_idx + 1}\n\n"
        try:
            # 加锁访问 PDF 对象，防止多线程竞争
            with self._pdf_lock:
                # 先获取表格边界框，用于后续从文本中排除表格区域
                table_bboxes = self._get_table_bboxes(page)
                # 提取文本内容（排除表格区域）
                text_content = self._extract_text_content(page, table_bboxes)
                # 提取表格内容
                tables_content = self._extract_tables(page, page_idx)
                # 收集页面中的图像数据
                raw_images = self._collect_page_raw_images(page)

            # 将提取的内容追加到页面 Markdown
            if text_content:
                page_md += text_content
            if tables_content:
                page_md += tables_content
            if raw_images:
                # 对图像进行 OCR 识别
                images_content = self._ocr_from_raw_images(raw_images, page_idx)
                if images_content:
                    page_md += images_content
        except Exception:
            # 页面处理异常，记录日志后继续处理下一页
            logger.error(f"第{page_idx + 1}页处理失败", exc_info=True)
        return page_md

    def extract_content(self, pdf_path: str) -> str:
        # 打开 PDF 文件
        pdf = pdfplumber.open(pdf_path)
        # 获取总页数
        total_pages = len(pdf.pages)
        # 获取所有页面对象
        pages = [pdf.pages[i] for i in range(total_pages)]

        try:
            # 存储各页处理结果
            results: Dict[int, str] = {}
            # 创建页面处理线程池
            with ThreadPoolExecutor(max_workers=PAGE_WORKERS) as executor:
                # 提交所有页面处理任务
                futures = {
                    executor.submit(self._process_single_page, pages[i], i): i
                    for i in range(total_pages)
                }
                # 收集完成的任务结果
                for future in as_completed(futures):
                    i = futures[future]
                    try:
                        results[i] = future.result()
                    except Exception:
                        # 处理失败时返回空页面
                        results[i] = f"## Page {i+1}\n\n"

            # 按页码顺序合并所有页面
            return "\n---\n\n".join(results[i] for i in range(total_pages))
        finally:
            pdf.close()

    # 探测 PDF 单页中所有表格的边界框（Bounding Box）坐标
    def _get_table_bboxes(self, page) -> List:
        try:
            tables = page.find_tables()
            return [t.bbox for t in tables]
        except Exception:
            return []

    # 提取 PDF 单页中所有文本内容（排除表格区域）中的文本内容
    def _extract_text_content(self, page, table_bboxes=None) -> str:
        # 提取页面文本内容（排除表格区域）
        text = self._extract_text_excluding_tables(page, table_bboxes)
        if not text or not text.strip():
            return ""

        # 初始化文本内容标题
        markdown = "### Text Content\n\n"
        # 按空行分割段落
        paragraphs = re.split(r'\n\s*\n', text.strip())
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            # 检测是否为键值对格式
            kv_match = self._detect_key_value_pair(para)
            if kv_match:
                key, value = kv_match
                # 键值对格式输出
                markdown += f"- **{key}**: {value}\n"
            else:
                # 普通段落格式输出
                markdown += f"{para}\n\n"
        return markdown + "\n"

    # 从PDF页面中提取纯文本，排除落在表格区域内的字符
    def _extract_text_excluding_tables(self, page, table_bboxes):
        if not table_bboxes:
            return page.extract_text() or ""

        chars = page.chars
        if not chars:
            return page.extract_text() or ""

        # 过滤掉落在表格区域内的字符
        filtered_chars = []
        for char in chars:
            cx = (char['x0'] + char['x1']) / 2
            cy = (char['top'] + char['bottom']) / 2
            in_table = any(
                bbox[0] <= cx <= bbox[2] and bbox[1] <= cy <= bbox[3]
                for bbox in table_bboxes
            )
            if not in_table:
                filtered_chars.append(char)
        if not filtered_chars:
            return ""

        return self._chars_to_text(filtered_chars)

    # 将 pdfplumber 的字符级对象列表，按阅读顺序重新拼成多行文本
    def _chars_to_text(self, chars):
        # 按 Y 坐标排序（同一行内按 X 坐标排序）
        chars = sorted(chars, key=lambda c: (round(c['top'], 1), c['x0']))

        lines = []
        current_line = []
        current_y = None

        for char in chars:
            char_y = round(char['top'], 1)
            if current_y is None:
                current_y = char_y
            # Y 坐标差异超过 3pt 视为新行
            if abs(char_y - current_y) > 3:
                if current_line:
                    lines.append(''.join(current_line))
                current_line = []
                current_y = char_y
            current_line.append(char['text'])

        if current_line:
            lines.append(''.join(current_line))

        return '\n'.join(lines)

    # 检测文本是否为键值对格式
    def _detect_key_value_pair(self, text: str) -> Optional[Tuple[str, str]]:
        # 定义多种键值对匹配模式
        patterns = [
            r'^([^:\n]{1,80}?)\s*[:：]\s*(.+)$',  # 冒号分隔（英文/日文）
            r'^([^-\n]{1,80}?)\s*-\s*(.+)$',     # 破折号分隔
            r'^([^=\n]{1,80}?)\s*=\s*(.+)$',     # 等号分隔
            r'^([^\s\n]{1,40})\s+(.+)$',         # 空格分隔
            r'^([^\(\n]+?)\((.+?)\)$',            # 括号形式
        ]
        for pattern in patterns:
            match = re.match(pattern, text.strip(), re.DOTALL)
            if match:
                key = match.group(1).strip()
                value = match.group(2).strip()
                # 验证键和值的有效性
                if (_is_valid_key(key) and
                        _is_valid_value(value) and
                        len(key) > 1 and len(value) > 0):
                    return key, value
        return None

    # 提取 PDF 单页中所有表格内容（排除文本区域）中的表格内容
    def _extract_tables(self, page, page_num: int) -> str:
        try:
            # 提取页面中的表格
            tables = page.extract_tables()
            if not tables:
                return ""
        except Exception:
            return ""

        # 初始化表格内容列表
        parts = ["### 表\n\n"]
        for idx, table in enumerate(tables):
            if not table or len(table) == 0:
                continue
            # 委托模块级共用函数转换表格为 Markdown
            table_md = _convert_table_to_md(table, idx)
            if table_md:
                parts.append(table_md + "\n\n")
        return "".join(parts)

    # 从 PDF 页面中提取所有嵌入图像（原始数据）
    def _collect_page_raw_images(self, page) -> List[Tuple[int, int, bytes]]:
        # 如果未启用 OCR，直接返回空列表
        if not self.use_ocr:
            return []
        try:
            # 获取页面中的嵌入图像
            embedded_imgs = page.images
            if not embedded_imgs:
                return []
        except Exception:
            return []

        raw_list: List[Tuple[int, int, bytes]] = []
        for img_info in embedded_imgs:
            try:
                # 获取图像原始尺寸
                w, h = img_info['srcsize']
                # 过滤过小的图像
                if w < 50 or h < 30:
                    continue
                # 获取图像原始数据
                raw_data = img_info['stream'].get_data()
                raw_list.append((w, h, raw_data))
            except Exception:
                continue
        return raw_list

    # 从原始图像数据进行 OCR 识别
    def _ocr_from_raw_images(
        self,
        raw_images: List[Tuple[int, int, bytes]],
        page_idx: int,
    ) -> str:
        # 阶段1：解码原始图像并提交 OCR 任务
        tasks = []
        for img_idx, (w, h, raw_data) in enumerate(raw_images):
            try:
                pil_img = None
                # 尝试不同的图像模式解码
                for mode, bpp in [('RGB', 3), ('L', 1), ('RGBA', 4), ('CMYK', 4)]:
                    if len(raw_data) == w * h * bpp:
                        try:
                            pil_img = Image.frombytes(mode, (w, h), raw_data)
                            if mode == 'CMYK':
                                pil_img = pil_img.convert('RGB')
                            break
                        except Exception:
                            continue
                if pil_img is None:
                    continue

                # 转换为灰度图
                if pil_img.mode == 'L':
                    gray = pil_img
                else:
                    gray = ImageOps.grayscale(pil_img)

                # 调整图像尺寸：过大图像影响 OCR 性能
                w, h = gray.size
                max_dim = max(w, h)
                if max_dim > MAX_OCR_IMG_PX:
                    scale = MAX_OCR_IMG_PX / max_dim
                    gray = gray.resize(
                        (int(w * scale), int(h * scale)), BILINEAR_FILTER
                    )

                # 提交 OCR 任务到线程池（直接使用模块级 _ocr_fast）
                future = self._ocr_executor.submit(
                    _ocr_fast, gray, SUPPORTED_LANGUAGES
                )
                tasks.append((img_idx, future, pil_img))
            except Exception:
                continue

        # 阶段2：收集 OCR 结果，必要时进行增强处理
        ocr_results = []
        for img_idx, future, pil_img in tasks:
            try:
                # 获取 OCR 结果，带超时
                ocr_text = future.result(timeout=OCR_TIMEOUT)
            except TimeoutError:
                continue

            # 如果 OCR 结果为空，尝试增强处理（直接使用模块级函数）
            if not ocr_text.strip():
                enhanced = _preprocess_fast(pil_img)
                # 调整增强后图像的尺寸
                w, h = enhanced.size
                max_dim = max(w, h)
                if max_dim > MAX_OCR_IMG_PX:
                    scale = MAX_OCR_IMG_PX / max_dim
                    enhanced = enhanced.resize(
                        (int(w * scale), int(h * scale)), BILINEAR_FILTER
                    )
                text2 = _ocr_fast(enhanced, SUPPORTED_LANGUAGES)
                if text2.strip():
                    ocr_text = text2

            ocr_text = ocr_text.strip()
            # 过滤过短的结果
            if ocr_text and len(ocr_text) > 5:
                cleaned = _clean_ocr_text(ocr_text)
                ocr_results.append(
                    f"### 図 {img_idx + 1}\n\n{cleaned}\n"
                )

        return "\n".join(ocr_results)

def parse_pdf(pdf_path: str, original_name: str) -> str:
    # 创建 PDF 转换器实例
    converter = PDFToMarkdownKVConverter(use_ocr=True)
    try:
        # 提取 PDF 内容
        raw = converter.extract_content(pdf_path)
        # 计算文件大小（MB）
        size_mb = os.path.getsize(pdf_path) / (1024 * 1024)
        # 添加元数据包装
        return wrap_markdown(original_name, raw, size_mb)
    finally:
        # 关闭 OCR 线程池，防止资源泄漏
        converter._ocr_executor.shutdown(wait=False)


# ================================================================
# 4. 画像解析
# ================================================================

# 解析图片文件（无页标记，走 flat 分片）
def parse_image(image_path: str, original_name: str) -> str:
    # 打开图像文件，失败则返回错误占位 Markdown
    try:
        pil_img = Image.open(image_path)
    except Exception:
        logger.error(f"无法打开图像文件: {image_path}", exc_info=True)
        return f"（画像解析失敗）\n"

    # 转为灰度图（复用共通 _convert_to_gray，减少色彩噪声对 OCR 的干扰）
    gray = _convert_to_gray(pil_img)

    # 智能缩放：长图按宽度限制，宽图按高度限制，避免文字被过度压缩
    w, h = gray.size
    aspect_ratio = h / w if w > 0 else 1.0
    if aspect_ratio > 2.5:
        # 长图（如网页长截图）：只限制宽度，高度保持原始分辨率
        if w > MAX_OCR_IMG_PX:
            scale = MAX_OCR_IMG_PX / w
            gray = gray.resize((int(w * scale), int(h * scale)), BILINEAR_FILTER)
    elif aspect_ratio < 0.4:
        # 宽图：只限制高度
        if h > MAX_OCR_IMG_PX:
            scale = MAX_OCR_IMG_PX / h
            gray = gray.resize((int(w * scale), int(h * scale)), BILINEAR_FILTER)
    else:
        # 普通图片：按最大边长等比缩放
        max_dim = max(w, h)
        if max_dim > MAX_OCR_IMG_PX:
            scale = MAX_OCR_IMG_PX / max_dim
            gray = gray.resize((int(w * scale), int(h * scale)), BILINEAR_FILTER)

    # 第一阶段 OCR：PSM 11（稀疏文本模式），失败则回退 PSM 3（全自动模式）
    text = _ocr_image_to_text(pil_img, gray)

    # 如果所有 OCR 尝试均无结果，返回占位 Markdown
    if not text.strip():
        return f"（文字認識できませんでした）\n"

    # 计算文件大小，用于元数据头部
    size_mb = os.path.getsize(image_path) / (1024 * 1024)

    # 将 OCR 文本作为 content 传入共通 wrap_markdown（无 ## Page N，走 flat 分片）
    content = f"{text}\n"
    return wrap_markdown(original_name, content, size_mb, doc_type="画像")


# ================================================================
# 5. DOCX 解析
# ================================================================

# 从 DOCX 文件中提取"图片引用ID（rId）→ 图片文件路径"的映射关系
def _build_docx_image_map(docx_path: str) -> Dict[str, str]:
    image_map: Dict[str, str] = {}
    try:
        with zipfile.ZipFile(docx_path, 'r') as zf:
            # 读取关系文件，建立 rId → target 映射
            rels_path = 'word/_rels/document.xml.rels'
            if rels_path not in zf.namelist():
                return image_map
            rels_xml = ET.parse(zf.open(rels_path))
            # 因为 XML 用了命名空间，找元素时必须加上
            ns = '{http://schemas.openxmlformats.org/package/2006/relationships}'
            for rel in rels_xml.findall(f'{ns}Relationship'):
                rel_id = rel.get('Id', '')
                target = rel.get('Target', '')
                # 只收集图片类型的关系
                if 'image' in target.lower() or target.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp')):
                    image_map[rel_id] = target
    except Exception:
        logger.warning("无法解析 DOCX 图片关系映射", exc_info=True)
    return image_map

# 从 DOCX 的一个段落（w:p）中，提取出里面所有纯文本内容，合并成一个干净的字符串返回。
def _extract_para_text(para_elem) -> str:
    # 存储的是 WordprocessingML 的主命名空间 URI
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    texts = []
    for t_elem in para_elem.findall(f'.//{ns}t'):
        if t_elem.text:
            texts.append(t_elem.text)
    return ''.join(texts).strip()

# 从 DOCX 段落元素中提取嵌入图片，进行 OCR 识别
def _extract_images_from_para(para_elem, image_rel_map: Dict[str, str],
                               docx_path: str) -> str:
    ns_w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    ns_drawing = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    ns_r = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

    ocr_texts: List[str] = []

    # 查找所有 w:drawing 元素（内联图片）
    for drawing in para_elem.findall(f'.//{ns_w}drawing'):
        # 在 drawing 内部查找 a:blip 元素（图片引用）
        for blip in drawing.findall(f'.//{ns_drawing}blip'):
            embed = blip.get(f'{ns_r}embed', '')
            if embed and embed in image_rel_map:
                # 获取图片在 zip 包中的路径
                img_name = image_rel_map[embed]
                # 从 zip 中提取图片并 OCR
                ocr_text = _ocr_docx_image(docx_path, img_name)
                if ocr_text:
                    ocr_texts.append(ocr_text)

    return '\n'.join(ocr_texts) if ocr_texts else ""

# 从 DOCX zip 包中提取指定图片，进行 OCR 识别
def _ocr_docx_image(docx_path: str, img_path_in_zip: str) -> str:
    try:
        with zipfile.ZipFile(docx_path, 'r') as zf:
            # 尝试多种路径格式
            raw_data = None
            for candidate in [img_path_in_zip,
                              f'word/media/{img_path_in_zip.split("/")[-1]}',
                              f'word/{img_path_in_zip.split("/")[-1]}']:
                if candidate in zf.namelist():
                    raw_data = zf.read(candidate)
                    break
            if raw_data is None:
                return ""

        # 从字节流打开图片
        pil_img = Image.open(io.BytesIO(raw_data))
        # 共通 OCR 管道：灰度 → 缩放 → 识别 → 清洗
        text = _ocr_image_from_pil(pil_img)
        return text if text.strip() else ""
    except Exception:
        logger.warning(f"DOCX 图片 OCR 失败: {img_path_in_zip}", exc_info=True)
        return ""

# 从 w:tbl 元素中提取表格数据（二维列表）
def _extract_docx_table(tbl_elem) -> List[List[str]]:
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    rows = tbl_elem.findall(f'{ns}tr')
    if not rows:
        return []
    table_data: List[List[str]] = []
    for row in rows:
        cells = row.findall(f'{ns}tc')
        row_data: List[str] = []
        for cell in cells:
            # 提取单元格内所有文本
            texts = []
            for p in cell.findall(f'{ns}p'):
                for t in p.findall(f'.//{ns}t'):
                    if t.text:
                        texts.append(t.text)
            row_data.append(''.join(texts).strip())
        # 跳过全空行
        if any(cell for cell in row_data):
            table_data.append(row_data)
    return table_data

# 解析 DOCX 文件（段落/表格/图片）
def parse_docx(path: str, original_name: str) -> str:
    try:
        import docx
        from docx.oxml.ns import qn
        # 打开 DOCX 文件
        doc = docx.Document(path)

        # 构建图片关系映射：rId → image_name（用于提取嵌入图片）
        image_rel_map = _build_docx_image_map(path)

        # 收集输出内容，按文档原始顺序
        parts: List[str] = []
        # 表格计数器
        table_idx = 0

        # 遍历 body 中的所有子元素，保持段落/表格/图片的原始顺序
        for child in doc.element.body:
            tag = child.tag

            if tag == qn('w:p'):
                # ── 段落处理 ──
                para_text = _extract_para_text(child)
                # 检查段落中是否包含图片
                img_text = _extract_images_from_para(child, image_rel_map, path)
                if img_text:
                    parts.append(img_text)
                if para_text:
                    parts.append(para_text)

            elif tag == qn('w:tbl'):
                # ── 表格处理 ──
                table_data = _extract_docx_table(child)
                if table_data:
                    table_md = _convert_table_to_md(table_data, table_idx)
                    if table_md:
                        parts.append(table_md)
                    table_idx += 1

        # 合并所有内容
        content = "\n\n".join(parts)
        if not content.strip():
            return ""

        # 计算文件大小，套用共通元数据头部
        size_mb = os.path.getsize(path) / (1024 * 1024)
        return wrap_markdown(original_name, content, size_mb, doc_type="DOCX")
    except Exception:
        logger.error(f"DOCX 解析失败: {path}", exc_info=True)
        return ""


# ================================================================
# 6. XLSX 解析
# ================================================================

# 从 .xlsx 中提取所有图片（按 sheet 索引）
def _extract_sheet_images(xlsx_path: str) -> Dict[int, str]:
    sheet_ocr_map: Dict[int, str] = {}

    try:
        with zipfile.ZipFile(xlsx_path, 'r') as zf:
            all_names = set(zf.namelist())

            # ---- 步骤1：建立 Sheet索引 → 图片文件名 映射 ----
            sheet_image_map: Dict[int, List[str]] = {}
            sheet_files = sorted([
                f for f in all_names
                if re.match(r'xl/worksheets/sheet\d+\.xml$', f)
            ])

            for idx, sheet_file in enumerate(sheet_files):
                sheet_num = os.path.basename(sheet_file).replace('sheet', '').replace('.xml', '')
                rels_file = f'xl/worksheets/_rels/sheet{sheet_num}.xml.rels'

                if rels_file not in all_names:
                    continue

                rels_tree = ET.parse(zf.open(rels_file))
                rels_root = rels_tree.getroot()

                drawing_rel_path = None
                for rel in rels_root:
                    if 'drawing' in rel.get('Type', '').lower():
                        target = rel.get('Target', '')
                        if target.startswith('../'):
                            drawing_rel_path = f'xl/{target[3:]}'
                        break

                if not drawing_rel_path or drawing_rel_path not in all_names:
                    continue

                drawing_dir = os.path.dirname(drawing_rel_path)
                drawing_name = os.path.basename(drawing_rel_path)
                drawing_rels_file = f'{drawing_dir}/_rels/{drawing_name}.rels'

                if drawing_rels_file not in all_names:
                    continue

                drawing_rels_tree = ET.parse(zf.open(drawing_rels_file))
                drawing_rels_root = drawing_rels_tree.getroot()

                images = []
                for rel in drawing_rels_root:
                    target = rel.get('Target', '')
                    img_name = target.replace('../media/', '').replace('media/', '')
                    images.append(img_name)

                if images:
                    sheet_image_map[idx] = images

            # ---- 步骤2：一次性遍历 xl/media/，OCR 所有图片 ----
            image_ocr_cache: Dict[str, str] = {}
            for name in all_names:
                if not name.startswith('xl/media/'):
                    continue
                if not name.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp')):
                    continue

                img_filename = name.replace('xl/media/', '')
                raw_data = zf.read(name)
                pil_img = Image.open(io.BytesIO(raw_data))
                # 共通 OCR 管道：灰度 → 缩放 → 识别 → 清洗
                text = _ocr_image_from_pil(pil_img)
                if text.strip():
                    image_ocr_cache[img_filename] = text

            # ---- 步骤3：按 Sheet 组装 OCR 文本 ----
            for sheet_idx, img_names in sheet_image_map.items():
                ocr_parts = [image_ocr_cache[n] for n in img_names if n in image_ocr_cache]
                if ocr_parts:
                    sheet_ocr_map[sheet_idx] = '\n'.join(ocr_parts)

    except Exception:
        logger.warning("Failed to extract XLSX sheet images", exc_info=True)

    return sheet_ocr_map

# 解析 XLSX 文件：Sheet → 表格/图片 → Markdown（按 ## Sheet: 分页）
def parse_xlsx(path: str, original_name: str) -> str:
    try:
        import openpyxl
        # 以只读模式打开 Excel，data_only=True 读取公式计算结果而非公式本身
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        parts = []
        table_idx = 0

        # 一次 ZIP 打开完成「Sheet→图片映射 + OCR」，返回 {sheet索引: "OCR文本"}
        sheet_ocr_map = _extract_sheet_images(path)

        # 遍历所有工作表
        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]

            # Sheet 标题（## Sheet: 模式，供 ChunkingService._split_by_sheet 识别）
            parts.append(f"## Sheet: {sheet_name}")

            # 当前 Sheet 的图片 OCR 文本
            img_ocr = sheet_ocr_map.get(sheet_idx, "")
            if img_ocr:
                parts.append(img_ocr)

            # 提取表格数据为二维列表
            table_data = []
            for row in ws.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                # 跳过全空行
                if any(cell for cell in row_data):
                    table_data.append(row_data)

            # 复用共通表格转换方法（PDF/DOCX/XLSX 共用）
            if table_data:
                table_md = _convert_table_to_md(table_data, table_idx)
                if table_md:
                    parts.append(table_md)
                table_idx += 1

        content = "\n\n".join(parts)
        if not content.strip():
            return ""

        # 计算文件大小，套用共通元数据头部
        size_mb = os.path.getsize(path) / (1024 * 1024)
        return wrap_markdown(original_name, content, size_mb, doc_type="XLSX")
    except Exception:
        logger.error(f"XLSX 解析失败: {path}", exc_info=True)
        return ""


# ================================================================
# 7. TXT 解析
# ================================================================

# 解析 TXT 文件（无页标记，走 flat 分片）
def parse_txt(path: str, original_name: str) -> str:
    try:
        # 导入 chardet 库用于编码检测
        import chardet
        # 读取文件二进制内容（使用 with 防止句柄泄漏）
        with open(path, "rb") as f:
            raw_bin = f.read()
        # 自动检测编码
        enc = chardet.detect(raw_bin)["encoding"] or "utf-8"
        # 解码为字符串
        text = raw_bin.decode(enc, errors="replace")
        # 计算文件大小，套用共通元数据头部（不添加 ## Page N，走 flat 分片路径）
        size_mb = os.path.getsize(path) / (1024 * 1024)
        return wrap_markdown(original_name, text, size_mb, doc_type="TXT")
    except Exception:
        logger.error(f"TXT 解析失败: {path}", exc_info=True)
        return ""